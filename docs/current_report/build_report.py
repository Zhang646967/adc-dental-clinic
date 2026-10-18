"""Build the current report from its reviewed JSON source using python-docx.

Run in a writable checkout with an explicitly selected Python environment.
PDF export/rendering is a separate verified LibreOffice step, documented in README.
No database, credentials, web service or host runtime is touched.
"""
from pathlib import Path
import argparse
import json
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def font(style, size):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr, value in [('ascii', 'Arial'), ('hAnsi', 'Arial'), ('eastAsia', 'PingFang SC')]:
        fonts.set(qn('w:' + attr), value)


def hyperlink(paragraph, label, url):
    from docx.opc.constants import RELATIONSHIP_TYPE
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79'); props.append(color)
    run.append(props)
    value = OxmlElement('w:t'); value.text = label; run.append(value)
    link.append(run); paragraph._p.append(link)


def table(doc, spec):
    result = doc.add_table(rows=1, cols=len(spec['headers']))
    result.autofit = False
    widths = spec['widths']
    for column, width in zip(result.columns, widths):column.width = Inches(width)
    for index, values in enumerate([spec['headers'], *spec['rows']]):
        row = result.rows[0] if index == 0 else result.add_row()
        properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); properties.append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); properties.append(repeat)
        for n, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[n]); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                line = OxmlElement('w:' + edge)
                for key, value in [('val','single'),('sz','4'),('color','D9D9D9')]:line.set(qn('w:' + key), value)
                borders.append(line)
            tcpr.append(borders)
            margin = OxmlElement('w:tcMar')
            for edge in ('top','bottom','left','right'):
                child = OxmlElement('w:' + edge); child.set(qn('w:w'), '95'); child.set(qn('w:type'), 'dxa'); margin.append(child)
            tcpr.append(margin)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '263B4B' if index == 0 else ('F0F4F6' if index % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.06
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if n in spec.get('center', []) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(text)); run.font.size = Pt(10.2)
            if index == 0:run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build(folder):
    source = json.loads((folder/'report_source.json').read_text(encoding='utf-8'))
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.70)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance = section.footer_distance = Inches(.30)
    for name, size in [('Normal',11),('Title',23),('Subtitle',12),('Heading 1',16),('Heading 2',12),('Caption',10)]:font(doc.styles[name],size)
    # The runtime template can carry decorative paragraph rules. Keep this report plain.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after = Pt(6); normal.line_spacing = 1.12; normal.widow_control = True
    for name in ('Heading 1','Heading 2'):
        p = doc.styles[name].paragraph_format
        p.space_before = Pt(9); p.space_after = Pt(7); p.keep_with_next = True
    doc.styles['Title'].paragraph_format.space_after = Pt(8)
    doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)
    head = section.header.paragraphs[0]; head.add_run('ADC  Systems Analysis and Design').font.size = Pt(8.5)
    foot = section.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('ADC  |  Page ').font.size = Pt(8.5)
    number = OxmlElement('w:fldSimple'); number.set(qn('w:instr'),'PAGE'); foot._p.append(number)
    doc.core_properties.title = source['title']
    doc.core_properties.subject = source['subtitle']
    doc.core_properties.author = 'ADC project'
    doc.core_properties.keywords = 'R1; R2 plan; synthetic course system; operational assurance'
    doc.add_paragraph(source['title'], 'Title')
    doc.add_paragraph(source['subtitle'], 'Subtitle')
    doc.add_paragraph(source['date_line'])
    markdown = ['# '+source['title'], '', source['subtitle'], '', source['date_line'], '']
    for index, item in enumerate(source['sections']):
        if index:doc.add_page_break()
        doc.add_heading(item['heading'],1)
        markdown += ['## '+item['heading'], '']
        for text in item.get('paragraphs',[]):doc.add_paragraph(text); markdown += [text,'']
        for spec in item.get('tables',[]):
            table(doc,spec)
            markdown += ['| '+' | '.join(spec['headers'])+' |', '| '+' | '.join('---' for _ in spec['headers'])+' |']
            markdown += ['| '+' | '.join(str(v).replace('|','/') for v in row)+' |' for row in spec['rows']]
            markdown.append('')
        for text in item.get('bullets',[]):doc.add_paragraph(text,'List Bullet'); markdown += ['- '+text]
        if item.get('bullets'):markdown.append('')
        for text in item.get('after',[]):doc.add_paragraph(text); markdown += [text,'']
        for reference in item.get('references',[]):
            p = doc.add_paragraph()
            p.add_run(reference['id']+' ')
            if reference.get('url'):hyperlink(p,reference['label'],reference['url'])
            else:p.add_run(reference['label'])
            if reference.get('detail'):p.add_run('. '+reference['detail'])
            label = '['+reference['label']+']('+reference['url']+')' if reference.get('url') else reference['label']
            markdown += [reference['id']+' '+label+'. '+reference.get('detail',''), '']
        if item.get('mind_map'):
            p = doc.add_paragraph(); p.add_run('Mind map 思维导图  ').bold = True; p.add_run(item['mind_map'])
            markdown += ['Mind map 思维导图: '+item['mind_map'],'']
    stem = source['file_stem']
    doc.save(folder/(stem+'.docx'))
    (folder/'REPORT.md').write_text('\n'.join(markdown).rstrip()+'\n',encoding='utf-8')
    print(json.dumps({'docx':stem+'.docx','sections':len(source['sections']),'markdown':'REPORT.md'}))


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--folder',type=Path,default=Path(__file__).resolve().parent)
    build(parser.parse_args().folder.resolve())
